# -*- coding: utf-8 -*-
"""TUTANAKLAR klasöründeki ilave iş tutanaklarını TAKİP TABLOSU.xlsx'e işler.

Kullanım:  python -X utf8 tutanak_gir.py          tabloya yazar
           python -X utf8 tutanak_gir.py --ozet   sadece okunanları listeler
Tabloda zaten bulunan tutanak numaraları atlanır (Sayfa1 ve Sayfa2 ayrı ayrı).
"""
import copy
import datetime
import glob
import os
import re
import sys

import openpyxl

TAKIP = "TAKİP TABLOSU.xlsx"
KLASOR = "TUTANAKLAR"
ANA_SAYFA = "İLAVE İS TUT."
SATIR_YUKSEKLIGI = 28.8


def kucuk(s):
    return str(s).replace("İ", "i").replace("I", "ı").lower()


def temiz(s):
    return re.sub(r"\s+", " ", str(s)).strip()


# ---------------------------------------------------------------- okuma
def bolum_satirlari(ws, harf):
    """'A)' / 'B)' başlığının altındaki dolu kalem satırlarını döndürür."""
    bas = next((r for r in range(30, ws.max_row + 1) if temiz(ws.cell(r, 2).value or "") == harf), None)
    if bas is None:
        return []
    satirlar = []
    for r in range(bas + 2, bas + 40):
        b = ws.cell(r, 2).value
        if isinstance(b, str) and ("toplam" in kucuk(b) or b.strip().endswith(")")):
            break
        if ws.cell(r, 3).value not in (None, ""):
            satirlar.append((temiz(ws.cell(r, 3).value), ws.cell(r, 6).value, ws.cell(r, 7).value))
    return satirlar


def rol_bul(ad):
    k = kucuk(ad)
    if "yard" in k:
        return "Usta Yard."
    if "usta" in k:
        return "Usta"
    if "işçi" in k:
        return "Düz İşçi"
    return None


def malzeme_duzenle(ad, birim, miktar):
    """Malzeme adını ve birimini tablodaki yazım biçimine çevirir."""
    k = kucuk(ad)
    b = {"m2": "m²", "m3": "m³", "mt": "m", "m": "m", "paket": "paket", "adet": "adet"}.get(
        kucuk(birim or "").strip(), birim)
    if re.match(r"\d+\.\d+\.\d+", k):                  # sözleşme pozu
        return (ad, b, miktar)
    if "gazbeton" in k and re.search(r"\d+", k):
        return (f"{re.search(r'\d+', k).group()} cm gazbeton", "m²", miktar)
    if "kutu profil" in k:
        return (re.sub(r"\s*mm\s*kutu profil", " mm Kutu Profil", ad, flags=re.I), "m", miktar)
    if "fx ultra" in k:
        return ("FX Ultra Alçıpan", "m²", miktar)
    if "suya dayanıklı" in k:
        return ("Suya Dayanıklı Alçıpan", "m²", miktar)
    if "mineral yün" in k:
        return ("4 cm Mineral Yün", "m²", miktar)
    return (ad, b, miktar)


def konum_ayir(metin):
    """'C3/5.20/C-1-067' -> (blok, kat, mahal)"""
    if re.match(r"C-[A-Z0-9]+-\d", str(metin or "").strip()):      # sadece mahal kodu yazılmış
        return None, None, str(metin).strip()
    p = [x.strip() for x in str(metin or "").split("/")]
    blok = p[0] if p and p[0] else None
    kat = p[1] if len(p) > 1 and p[1] else None
    mahal = "/".join(p[2:]) if len(p) > 2 else None
    return blok, kat, mahal


def tutanak_oku(dosya):
    wb = openpyxl.load_workbook(dosya, data_only=True)
    ws = wb[ANA_SAYFA]
    ad = os.path.basename(dosya)
    uyari = []

    no_ham = str(ws["D7"].value or "").strip()
    m = re.search(r"AAS-(\d+)([A-Z]?)", ad)
    no = int(m.group(1))
    ek = m.group(2)
    if not re.fullmatch(r"0*%d%s" % (no, ek), no_ham.replace(" ", ""), flags=re.I):
        uyari.append(f"dosya adı {no}{ek}, içindeki no '{no_ham}'")

    tarih = ws["D5"].value
    if not isinstance(tarih, datetime.datetime):
        uyari.append(f"tarih okunamadı: {tarih!r}")

    aciklama = " ".join(temiz(ws.cell(r, 5).value) for r in range(28, 32)
                        if ws.cell(r, 5).value and not str(ws.cell(r, 5).value).startswith("*"))
    if not aciklama:
        uyari.append("açıklama boş")

    # A) bölümü: işçilik + (varsa) poz/malzeme, B) bölümü: REC malzemesi.
    # Bazı tutanaklarda işçiler B bölümüne de yazılmış; ada göre ayrılıyor.
    kisiler, malzeme = [], []
    for harf in ("A)", "B)"):
        for kalem, birim, miktar in bolum_satirlari(ws, harf):
            b = kucuk(birim or "")
            rol = rol_bul(kalem)
            if rol or "yevmiye" in b or "adam" in b:
                if rol is None and "ay" in b:
                    rol = "İSG Personeli"
                if rol is None:
                    uyari.append(f"'{kalem}' usta/yardımcı belirtilmemiş, Usta sayıldı")
                    rol = "Usta"
                if not isinstance(miktar, (int, float)):
                    uyari.append(f"'{kalem}' yevmiye miktarı yok")
                    miktar = 0
                kisiler.append((kalem, rol, "ücret" if "ay" in b else "yevmiye", miktar))
            else:
                if not isinstance(miktar, (int, float)):
                    uyari.append(f"'{kalem}' miktarı yok")
                if isinstance(miktar, float):
                    miktar = round(miktar, 4)
                malzeme.append(malzeme_duzenle(kalem, birim, miktar))

    return {
        "no": no, "ek": ek, "etiket": f"{no}{ek}", "dosya": ad, "tarih": tarih,
        "konum": konum_ayir(ws["I9"].value), "aciklama": aciklama,
        "kisiler": kisiler, "malzeme": malzeme,
        "iptal": "iptal" in kucuk(ad), "uyari": uyari,
    }


def iscilik_satirlari(t):
    """Sayfa2 için (yevmiye detayı, birim, miktar) satırları."""
    if any(r == "Düz İşçi" for _, r, _, _ in t["kisiler"]):     # İSG: kişi kişi
        return [(re.sub(r"\s*\(", " (", k), b, m) for k, _, b, m in t["kisiler"]]
    out = []
    for rol in ("Usta", "Usta Yard.", "İSG Personeli"):
        grup = [(b, m) for _, r, b, m in t["kisiler"] if r == rol]
        if grup:
            out.append((f"{len(grup)} {rol}", grup[0][0], sum(m for _, m in grup)))
    return out


def tumunu_oku():
    hepsi = {}
    for d in sorted(glob.glob(os.path.join(KLASOR, "*.xlsx"))):
        if os.path.basename(d).startswith("~$"):
            continue
        try:
            t = tutanak_oku(d)
        except Exception as e:                       # farklı şablon vb.
            print(f"OKUNAMADI  {os.path.basename(d)}: {e}")
            continue
        onceki = hepsi.get(t["etiket"])
        if onceki and "rev" not in kucuk(t["dosya"]):
            continue
        hepsi[t["etiket"]] = t
    return sorted(hepsi.values(), key=lambda t: (t["no"], t["ek"]))


# ---------------------------------------------------------------- yazma
def son_dolu_satir(ws, c1, c2):
    for r in range(ws.max_row, 3, -1):
        if any(ws.cell(r, c).value is not None for c in range(c1, c2 + 1)):
            return r
    return 3


def mevcut_nolar(ws, sutun):
    return {str(c.value).strip() for c in ws[sutun][3:] if c.value is not None and not str(c.value).startswith("=")}


def blok_hazirla(ws, ilk, son, ornek, c1, c2, birlesen):
    for m in [m for m in ws.merged_cells.ranges if m.min_row <= son and m.max_row >= ilk]:
        ws.unmerge_cells(str(m))
    for r in range(ilk, son + 1):
        ws.row_dimensions[r].height = SATIR_YUKSEKLIGI
        for c in range(c1, c2 + 1):
            ws.cell(r, c).value = None
            ws.cell(r, c)._style = copy.copy(ws.cell(ornek, c)._style)
    if son > ilk:
        for c in birlesen:
            ws.merge_cells(start_row=ilk, start_column=c, end_row=son, end_column=c)


def sayfa2_yaz(ws, t, ornek):
    isc = iscilik_satirlari(t)
    ilk = son_dolu_satir(ws, 2, 12) + 1
    son = ilk + max(len(isc), len(t["malzeme"]), 1) - 1
    blok_hazirla(ws, ilk, son, ornek, 2, 16, range(2, 7))
    if t["iptal"]:
        ws.cell(ilk, 1, "İPTAL")
    ws.cell(ilk, 2, t["tarih"])
    ws.cell(ilk, 3, t["no"] if not t["ek"] else t["etiket"])
    ws.cell(ilk, 4, t["aciklama"])
    if isc:
        ws.cell(ilk, 5, "\n".join(e for e, _, _ in isc))
        ws.cell(ilk, 6, isc[0][1])
    for i, (_, _, miktar) in enumerate(isc):
        ws.cell(ilk + i, 7, miktar)
        if not t["iptal"]:
            ws.cell(ilk + i, 9, f"=+G{ilk + i}*H{ilk + i}")
    for i, (ad, birim, miktar) in enumerate(t["malzeme"]):
        ws.cell(ilk + i, 10, ad)
        ws.cell(ilk + i, 11, birim)
        ws.cell(ilk + i, 12, miktar)


def sayfa1_yaz(ws, t, ornek):
    ilk = son_dolu_satir(ws, 3, 23) + 1
    son = ilk + max(len(t["malzeme"]), 1) - 1
    blok_hazirla(ws, ilk, son, ornek, 2, 24, range(2, 16))
    if t["iptal"]:
        ws.cell(ilk, 1, "İPTAL")
    ws.cell(ilk, 3, t["etiket"])
    ws.cell(ilk, 4, t["tarih"])
    for c, v in zip((5, 6, 7), t["konum"]):
        ws.cell(ilk, c, v)
    usta = sum(m for _, r, _, m in t["kisiler"] if r == "Usta")
    yard = sum(m for _, r, _, m in t["kisiler"] if r in ("Usta Yard.", "Düz İşçi"))
    ws.cell(ilk, 8, usta or None)
    ws.cell(ilk, 10, yard or None)
    for i, (ad, birim, miktar) in enumerate(t["malzeme"]):
        ws.cell(ilk + i, 16, ad)
        ws.cell(ilk + i, 17, miktar)
        ws.cell(ilk + i, 18, birim)


def main():
    tutanaklar = tumunu_oku()
    if "--ozet" in sys.argv:
        for t in tutanaklar:
            print(t["etiket"], t["tarih"].strftime("%d.%m.%y") if hasattr(t["tarih"], "strftime") else t["tarih"],
                  "|".join(str(x) for x in t["konum"]), "İPTAL" if t["iptal"] else "",
                  iscilik_satirlari(t), t["malzeme"], len(t["aciklama"]), "!! " + "; ".join(t["uyari"]) if t["uyari"] else "")
        return
    wb = openpyxl.load_workbook(TAKIP)
    s1, s2 = wb["Sayfa1"], wb["Sayfa2"]
    var1, var2 = mevcut_nolar(s1, "C"), mevcut_nolar(s2, "C")
    ornek1, ornek2 = son_dolu_satir(s1, 3, 23), son_dolu_satir(s2, 2, 12)
    n1 = n2 = 0
    for t in tutanaklar:
        if t["etiket"] not in var1:
            sayfa1_yaz(s1, t, ornek1); n1 += 1
        if t["etiket"] not in var2:
            sayfa2_yaz(s2, t, ornek2); n2 += 1
    s1["X4"] = f"=SUM(X5:X{son_dolu_satir(s1, 3, 23)})"
    wb.save(TAKIP)
    print(f"Sayfa1: {n1} tutanak eklendi, Sayfa2: {n2} tutanak eklendi")


if __name__ == "__main__":
    main()
